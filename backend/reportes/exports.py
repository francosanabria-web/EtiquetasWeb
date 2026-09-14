# -*- coding: utf-8 -*-
"""Export Excel (.xlsx) en formato Tabla (openpyxl Table).

Columnas = detalle del mail diario de gastos (escritorio) + SECTOR
para listado unificado.
"""

from __future__ import annotations

import io
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from config import COLUMNAS_EXPORT_TABLA

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_MONEY_COLS = frozenset({"PRECIO_UNITARIO", "MONTO_TOTAL_SALIDA"})


def _header_display(col: str) -> str:
    """Igual que el mail HTML: guiones bajos → espacios."""
    return str(col).replace("_", " ")


def df_detalle_gastos(df: pd.DataFrame) -> pd.DataFrame:
    """Proyecta al layout del detalle de retiros del mail (+ SECTOR)."""
    if df is None or df.empty:
        return pd.DataFrame(columns=COLUMNAS_EXPORT_TABLA)
    out = df.copy()
    for col in COLUMNAS_EXPORT_TABLA:
        if col not in out.columns:
            out[col] = "" if col not in ("CANTIDAD", "PRECIO_UNITARIO", "MONTO_TOTAL_SALIDA") else 0.0
    return out[COLUMNAS_EXPORT_TABLA].copy()


def build_xlsx_tabla(
    df: pd.DataFrame,
    *,
    sheet_name: str = "Movimientos",
    table_name: str = "TablaGastos",
) -> bytes:
    """Genera .xlsx con una Excel Table (estilo TableStyleMedium2)."""
    data = df_detalle_gastos(df)
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Movimientos")[:31]

    headers = [_header_display(c) for c in COLUMNAS_EXPORT_TABLA]
    for col_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _THIN

    for row_idx, row in enumerate(data.itertuples(index=False), start=2):
        for col_idx, (col_name, val) in enumerate(zip(COLUMNAS_EXPORT_TABLA, row), start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = _THIN
            if col_name in _MONEY_COLS or col_name == "CANTIDAD":
                try:
                    num = float(val) if val is not None and str(val).strip() != "" else 0.0
                except (TypeError, ValueError):
                    num = 0.0
                cell.value = num
                if col_name in _MONEY_COLS:
                    cell.number_format = '"$"#,##0.00'
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.value = "" if val is None or (isinstance(val, float) and pd.isna(val)) else str(val)

    n_rows = max(1, len(data) + 1)  # header + datos (mín. 1 fila de header)
    n_cols = len(COLUMNAS_EXPORT_TABLA)
    # Excel Table requiere al menos header; si no hay datos, una fila vacía
    if len(data) == 0:
        for col_idx in range(1, n_cols + 1):
            ws.cell(row=2, column=col_idx, value="").border = _THIN
        n_rows = 2

    ref = f"A1:{get_column_letter(n_cols)}{n_rows}"
    table = Table(displayName=table_name[:255], ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    for col_idx, col_name in enumerate(COLUMNAS_EXPORT_TABLA, start=1):
        width = min(max(12, len(_header_display(col_name)) + 4), 40)
        if col_name == "DESCRIPCION":
            width = 36
        elif col_name in _MONEY_COLS:
            width = 16
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ref

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def meta_columnas_export() -> list[dict[str, Any]]:
    return [
        {"campo": c, "titulo": _header_display(c)}
        for c in COLUMNAS_EXPORT_TABLA
    ]
