# -*- coding: utf-8 -*-
"""Export PDF / Excel de activos (fuera o ingresados)."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any, Literal

from config import DIAS_ALERTA, DIAS_AVISO
from service import activos_resumen
from store import ActivosStore

Vista = Literal["fuera", "ingresados"]

_VERDE = "C6EFCE"
_AMARILLO = "FFEB9C"
_ROJO = "FFC7CE"

_COLUMNAS_FUERA = [
    ("dias_fuera", "Días fuera"),
    ("equipo", "Equipo / repuesto"),
    ("cantidad", "Cantidad"),
    ("observaciones", "Observaciones"),
    ("proveedor", "Proveedor"),
    ("sector", "Sector"),
]

_COLUMNAS_ING = [
    ("dias_fuera", "Días fuera"),
    ("equipo", "Equipo / repuesto"),
    ("cantidad", "Cantidad"),
    ("observaciones", "Observaciones"),
    ("proveedor", "Proveedor"),
    ("sector", "Sector"),
]


def _color_dias(dias: int) -> str:
    if dias > DIAS_ALERTA:
        return _ROJO
    if dias > DIAS_AVISO:
        return _AMARILLO
    return _VERDE


def _estado_legible(s: str) -> str:
    return str(s or "").replace("_", " ").strip() or "—"


def _lista_y_meta(store: ActivosStore, vista: Vista) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    data = activos_resumen(store)
    if vista == "ingresados":
        return data.get("lista_ingresados", []), {
            "titulo": "Activos ingresados a planta",
            "hoja": "Ingresados",
            "archivo": "activos_ingresados",
            "subtitulo": (
                f"{data.get('ingresados', 0)} ingresados · "
                f"generado {datetime.now().strftime('%d/%m/%Y %H:%M')}"
            ),
            "columnas": _COLUMNAS_ING,
            "color_header": "1F4E78",
            "colorear_dias": False,
        }
    return data.get("lista_fuera", []), {
        "titulo": "Activos fuera de planta",
        "hoja": "Fuera de planta",
        "archivo": "activos_fuera_de_planta",
        "subtitulo": (
            f"{data.get('fuera_de_planta', 0)} equipos fuera · "
            f"promedio {data.get('dias_promedio_fuera', 0)} días · "
            f"generado {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ),
        "columnas": _COLUMNAS_FUERA,
        "color_header": "A32D2D",
        "colorear_dias": True,
    }


def activos_excel(store: ActivosStore, vista: Vista = "fuera") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    lista, meta = _lista_y_meta(store, vista)
    columnas = meta["columnas"]

    wb = Workbook()
    ws = wb.active
    ws.title = meta["hoja"]
    ws.sheet_view.showGridLines = False

    n_cols = len(columnas)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    t = ws.cell(row=1, column=1, value=meta["titulo"])
    t.font = Font(bold=True, size=15, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=meta["color_header"])
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    sub = ws.cell(row=2, column=1, value=meta["subtitulo"])
    sub.font = Font(size=10, italic=True, color="555555")
    sub.alignment = Alignment(horizontal="left", indent=1)

    thin = Side(style="thin", color="BFBFBF")
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)

    fila = 4
    for col, (_key, titulo) in enumerate(columnas, start=1):
        c = ws.cell(row=fila, column=col, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="475569")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = borde
    ws.freeze_panes = f"A{fila + 1}"

    for it in lista:
        fila += 1
        dias = int(it.get("dias_fuera", 0) or 0)
        color = _color_dias(dias) if meta["colorear_dias"] else "FFFFFF"
        for col, (key, _titulo) in enumerate(columnas, start=1):
            if key in ("estado", "estado_al_ingreso"):
                val: Any = _estado_legible(str(it.get(key, "") or ""))
            elif key == "dias_fuera":
                val = dias
            elif key == "cantidad":
                try:
                    v = it.get(key)
                    val = int(float(v)) if v not in (None, "", "—") else 1
                except Exception:
                    val = it.get(key) or "—"
            else:
                val = it.get(key) or "—"
            c = ws.cell(row=fila, column=col, value=val)
            if meta["colorear_dias"]:
                c.fill = PatternFill("solid", fgColor=color)
            c.border = borde
            c.alignment = Alignment(
                horizontal="center" if key in ("dias_fuera", "cantidad", "sector") else "left",
                vertical="center",
            )

    for i in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    # Equipo / repuesto is now column B (dias_fuera is A) — give it more width
    if n_cols >= 2:
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 34
    else:
        ws.column_dimensions["A"].width = 32

    out = BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def activos_pdf(store: ActivosStore, vista: Vista = "fuera") -> bytes:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    lista, meta = _lista_y_meta(store, vista)
    columnas = meta["columnas"]

    def s(v: Any, n: int | None = None) -> str:
        txt = str(v if v not in (None, "") else "-")
        txt = txt.encode("latin-1", "replace").decode("latin-1")
        if n and len(txt) > n:
            return txt[: n - 1] + "."
        return txt

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(0, 9, s(meta["titulo"]), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 6, s(meta["subtitulo"]), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(2)

    # Anchos adaptados a cantidad de columnas — priorizar Equipo / repuesto (col 1 = Días, col 2 = Equipo)
    n = len(columnas)
    usable = 277
    widths = [usable / n] * n
    if n >= 2:
        # Equipo es columna 1 (índice 1) tras Días fuera
        widths[1] = min(78, usable * 0.30)
        widths[0] = min(22, usable * 0.10)  # Días fuera angosto
        resto = usable - widths[0] - widths[1]
        otros = n - 2
        if otros > 0:
            for i in range(2, n):
                widths[i] = resto / otros
    elif n >= 1:
        widths[0] = min(58, usable * 0.22)
        resto = usable - widths[0]
        for i in range(1, n):
            widths[i] = resto / (n - 1)

    pdf.set_font("Helvetica", "B", 7)
    pdf.set_fill_color(71, 85, 105)
    pdf.set_text_color(255, 255, 255)
    for w, (_k, tit) in zip(widths, columnas):
        pdf.cell(w, 7, s(tit, 22), border=1, fill=True, align="C", new_x=XPos.RIGHT, new_y=YPos.TOP)
    pdf.ln()

    pdf.set_font("Helvetica", "", 7)
    pdf.set_text_color(20, 20, 20)
    hex_rgb = {
        _VERDE: (198, 239, 206),
        _AMARILLO: (255, 235, 156),
        _ROJO: (255, 199, 206),
        "FFFFFF": (255, 255, 255),
    }
    for it in lista:
        dias = int(it.get("dias_fuera", 0) or 0)
        color = _color_dias(dias) if meta["colorear_dias"] else "FFFFFF"
        r, g, b = hex_rgb[color]
        pdf.set_fill_color(r, g, b)
        for w, (key, _t) in zip(widths, columnas):
            if key in ("estado", "estado_al_ingreso"):
                val = s(_estado_legible(str(it.get(key, "") or "")), 18)
            elif key == "dias_fuera":
                val = str(dias)
            elif key == "cantidad":
                try:
                    v = it.get(key)
                    val = str(int(float(v))) if v not in (None, "", "—") else "1"
                except Exception:
                    val = s(it.get(key), 8)
            elif key == "observaciones":
                val = s(it.get(key), 32)
            else:
                val = s(it.get(key), 28 if key == "equipo" else 14)
            align = "C" if key in ("dias_fuera", "cantidad", "sector") else "L"
            pdf.cell(w, 6, val, border=1, fill=True, align=align, new_x=XPos.RIGHT, new_y=YPos.TOP)
        pdf.ln()

    out = BytesIO()
    pdf.output(out)
    return out.getvalue()


def nombre_archivo(vista: Vista, ext: str) -> str:
    base = "activos_ingresados" if vista == "ingresados" else "activos_fuera_de_planta"
    return f"{base}.{ext}"
