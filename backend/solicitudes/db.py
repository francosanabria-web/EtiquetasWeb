# -*- coding: utf-8 -*-
"""Persistencia — Solicitud de pedidos.

Fuente de verdad: un único Excel en la carpeta compartida de Google Drive
(`config.EXCEL_PATH`). Todas las PC leen/escriben ese archivo, igual que la app
de escritorio, así los datos no viven en una base local que se pueda perder al
actualizar el programa. Más adelante se migra a MariaDB/HeidiSQL.

Estructura del Excel:
- Hoja ``PEDIDOS``: una fila por ítem (la cabecera del pedido se repite en cada
  ítem del mismo pedido). Se agrupa por ``SOLICITUD_ID`` al leer.
- Hoja ``CONTADORES``: pares clave/valor para los correlativos
  (``pedido`` → P-XXXX, ``tr`` → TR-XXXX, ``id_seq`` → id interno).

Escrituras: bajo un lock de archivo (coordina varias PC sobre Drive) + escritura
atómica (archivo temporal + ``os.replace``) + copia ``.bak`` previa.
"""

from __future__ import annotations

import os
import sqlite3
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import DB_PATH, EXCEL_PATH, UPLOADS_DIR, cargar_catalogos

SHEET_PEDIDOS = "PEDIDOS"
SHEET_CONTADORES = "CONTADORES"

# Encabezado canónico de la hoja PEDIDOS (una fila por ítem).
COLUMNS: list[tuple[str, str]] = [
    ("SOLICITUD_ID", "solicitud_id"),
    ("N_PEDIDO", "n_pedido"),
    ("N_TR", "n_tr"),
    ("TIPO", "tipo"),
    ("ESTADO", "estado"),
    ("CUENTA_CONTABLE", "cuenta_contable"),
    ("SOLICITANTE", "solicitante"),
    ("PROVEEDOR", "proveedor"),
    ("REMITO_NRO", "remito_nro"),
    ("REMITO_ARCHIVO", "remito_archivo"),
    ("PRESUPUESTO_NRO", "presupuesto_nro"),
    ("PRESUPUESTO_ARCHIVO", "presupuesto_archivo"),
    ("NOTAS", "notas"),
    ("ITEM_ORDEN", "orden"),
    ("ITEM_CODIGO", "codigo"),
    ("ITEM_DESCRIPCION", "descripcion"),
    ("ITEM_CANTIDAD", "cantidad"),
    ("ITEM_UNIDAD", "unidad"),
    ("ITEM_AREA_MAQUINA", "area"),
    ("ITEM_IMAGEN", "imagen_path"),
    ("CREADO_POR", "creado_por"),
    ("CREADO_EN", "creado_en"),
    ("ACTUALIZADO_EN", "actualizado_en"),
]
HEADER_TO_KEY = {h: k for h, k in COLUMNS}
HEADERS = [h for h, _ in COLUMNS]

# Campos que pertenecen a la cabecera del pedido (se repiten en cada ítem).
CABECERA_KEYS = [
    "n_pedido",
    "n_tr",
    "tipo",
    "estado",
    "cuenta_contable",
    "solicitante",
    "proveedor",
    "remito_nro",
    "remito_archivo",
    "presupuesto_nro",
    "presupuesto_archivo",
    "notas",
    "creado_por",
    "creado_en",
    "actualizado_en",
]
ITEM_KEYS = ["orden", "codigo", "descripcion", "cantidad", "unidad", "area", "imagen_path"]

_LOCK_TIMEOUT_S = 12.0
_LOCK_STALE_S = 60.0

# Cache en memoria (evita releer el Excel en cada request de solo lectura).
_cache: dict[str, Any] = {"mtime": None, "solicitudes": None, "contadores": None}


def _ahora_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


# --------------------------------------------------------------------------- #
# Lock de archivo (coordinación multi-PC sobre Drive)
# --------------------------------------------------------------------------- #
class _FileLock:
    def __init__(self, target: Path) -> None:
        self.lock_path = target.with_suffix(target.suffix + ".lock")
        self._fd: int | None = None

    def __enter__(self) -> "_FileLock":
        inicio = time.time()
        while True:
            try:
                self._fd = os.open(
                    str(self.lock_path), os.O_CREAT | os.O_EXCL | os.O_RDWR
                )
                os.write(self._fd, str(os.getpid()).encode("ascii", "ignore"))
                return self
            except FileExistsError:
                # ¿Lock viejo/olvidado? Robarlo pasado el umbral.
                try:
                    edad = time.time() - self.lock_path.stat().st_mtime
                    if edad > _LOCK_STALE_S:
                        self.lock_path.unlink(missing_ok=True)
                        continue
                except OSError:
                    pass
                if time.time() - inicio > _LOCK_TIMEOUT_S:
                    raise TimeoutError(
                        "El archivo de pedidos está ocupado por otra PC. Reintentá en unos segundos."
                    )
                time.sleep(0.15)

    def __exit__(self, *_exc: object) -> None:
        try:
            if self._fd is not None:
                os.close(self._fd)
        finally:
            self.lock_path.unlink(missing_ok=True)


# --------------------------------------------------------------------------- #
# Lectura
# --------------------------------------------------------------------------- #
def _asegurar_dirs() -> None:
    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    UPLOADS_DIR.mkdir(parents=True, exist_ok=True)


def _celda_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _leer_workbook() -> tuple[list[dict[str, Any]], dict[str, int]]:
    """Lee el Excel y reconstruye pedidos (agrupando ítems) + contadores."""
    if not EXCEL_PATH.is_file():
        return [], {}

    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    solicitudes: dict[int, dict[str, Any]] = {}
    orden_aparicion: list[int] = []

    if SHEET_PEDIDOS in wb.sheetnames:
        ws = wb[SHEET_PEDIDOS]
        filas = ws.iter_rows(values_only=True)
        try:
            cabecera = [str(c or "").strip() for c in next(filas)]
        except StopIteration:
            cabecera = []
        idx = {h: i for i, h in enumerate(cabecera)}

        def val(row: tuple[Any, ...], header: str) -> str:
            i = idx.get(header)
            if i is None or i >= len(row):
                return ""
            return _celda_str(row[i])

        for row in filas:
            if row is None or all(c is None for c in row):
                continue
            sid_raw = val(row, "SOLICITUD_ID")
            if not sid_raw:
                continue
            try:
                sid = int(float(sid_raw))
            except ValueError:
                continue

            if sid not in solicitudes:
                cab: dict[str, Any] = {"id": sid}
                for header, key in COLUMNS:
                    if key in CABECERA_KEYS:
                        cab[key] = val(row, header)
                cab["items"] = []
                solicitudes[sid] = cab
                orden_aparicion.append(sid)

            desc = val(row, "ITEM_DESCRIPCION")
            cod = val(row, "ITEM_CODIGO")
            if desc or cod:
                try:
                    cant = float(val(row, "ITEM_CANTIDAD") or 0)
                except ValueError:
                    cant = 0.0
                try:
                    orden = int(float(val(row, "ITEM_ORDEN") or 0))
                except ValueError:
                    orden = len(solicitudes[sid]["items"])
                solicitudes[sid]["items"].append(
                    {
                        "orden": orden,
                        "codigo": cod,
                        "descripcion": desc,
                        "cantidad": cant,
                        "unidad": val(row, "ITEM_UNIDAD"),
                        "area": val(row, "ITEM_AREA_MAQUINA"),
                        "imagen_path": val(row, "ITEM_IMAGEN"),
                    }
                )

    contadores: dict[str, int] = {}
    if SHEET_CONTADORES in wb.sheetnames:
        wsc = wb[SHEET_CONTADORES]
        for row in wsc.iter_rows(min_row=2, values_only=True):
            if not row or row[0] in (None, ""):
                continue
            clave = str(row[0]).strip().lower()
            try:
                contadores[clave] = int(float(row[1] or 0))
            except (ValueError, TypeError, IndexError):
                contadores[clave] = 0

    wb.close()

    lista = []
    for sid in orden_aparicion:
        sol = solicitudes[sid]
        sol["items"].sort(key=lambda it: (int(it.get("orden") or 0)))
        lista.append(sol)
    return lista, contadores


def _load_all(*, force: bool = False) -> tuple[list[dict[str, Any]], dict[str, int]]:
    """Devuelve (solicitudes, contadores) usando cache por mtime."""
    _asegurar_dirs()
    mtime = EXCEL_PATH.stat().st_mtime if EXCEL_PATH.is_file() else None
    if (
        not force
        and _cache["solicitudes"] is not None
        and _cache["mtime"] == mtime
    ):
        return _cache["solicitudes"], _cache["contadores"]
    solicitudes, contadores = _leer_workbook()
    _cache["mtime"] = mtime
    _cache["solicitudes"] = solicitudes
    _cache["contadores"] = contadores
    return solicitudes, contadores


# --------------------------------------------------------------------------- #
# Escritura atómica
# --------------------------------------------------------------------------- #
def _escribir_workbook(solicitudes: list[dict[str, Any]], contadores: dict[str, int]) -> None:
    _asegurar_dirs()
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_PEDIDOS

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    ws.append(HEADERS)
    for col in range(1, len(HEADERS) + 1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for sol in solicitudes:
        items = sol.get("items") or [{}]
        for it in items:
            fila = []
            for header, key in COLUMNS:
                if header == "SOLICITUD_ID":
                    fila.append(int(sol["id"]))
                elif key in ITEM_KEYS:
                    v = it.get(key, "")
                    if key == "cantidad":
                        try:
                            v = float(v or 0)
                        except (ValueError, TypeError):
                            v = 0
                    elif key == "orden":
                        try:
                            v = int(float(v or 0))
                        except (ValueError, TypeError):
                            v = 0
                    fila.append(v)
                else:
                    fila.append(sol.get(key, ""))
            ws.append(fila)

    anchos = {
        "ITEM_DESCRIPCION": 38,
        "NOTAS": 30,
        "CUENTA_CONTABLE": 32,
        "PROVEEDOR": 24,
        "SOLICITANTE": 22,
        "ITEM_AREA_MAQUINA": 22,
        "CREADO_EN": 22,
        "ACTUALIZADO_EN": 22,
    }
    for i, (header, _key) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = anchos.get(header, 14)

    wsc = wb.create_sheet(SHEET_CONTADORES)
    wsc.append(["CLAVE", "VALOR"])
    for col in range(1, 3):
        wsc.cell(row=1, column=col).font = Font(bold=True)
    for clave in sorted(contadores):
        wsc.append([clave, int(contadores[clave])])

    tmp = EXCEL_PATH.with_suffix(EXCEL_PATH.suffix + ".tmp")
    wb.save(tmp)
    wb.close()

    if EXCEL_PATH.is_file():
        try:
            bak = EXCEL_PATH.with_suffix(EXCEL_PATH.suffix + ".bak")
            os.replace(EXCEL_PATH, bak)
        except OSError:
            pass
    os.replace(tmp, EXCEL_PATH)

    # Refrescar cache con lo recién escrito.
    _cache["mtime"] = EXCEL_PATH.stat().st_mtime
    _cache["solicitudes"] = solicitudes
    _cache["contadores"] = contadores


# --------------------------------------------------------------------------- #
# Contadores / correlativos
# --------------------------------------------------------------------------- #
def _next(contadores: dict[str, int], clave: str) -> int:
    contadores[clave] = int(contadores.get(clave, 0)) + 1
    return contadores[clave]


# --------------------------------------------------------------------------- #
# Enriquecido / validación
# --------------------------------------------------------------------------- #
def _enrich(sol: dict[str, Any]) -> dict[str, Any]:
    items = sol.get("items") or []
    d = dict(sol)
    # Normalizar cabecera para que nunca falte una clave esperada por el front.
    for key in CABECERA_KEYS:
        d.setdefault(key, "")
    d["items"] = items
    d["items_count"] = len(items)
    if items:
        primero = items[0]
        d["preview"] = primero.get("descripcion") or primero.get("codigo") or ""
        areas = sorted(
            {str(i.get("area") or "").strip() for i in items if str(i.get("area") or "").strip()}
        )
        d["areas_resumen"] = ", ".join(areas)
    else:
        d["preview"] = ""
        d["areas_resumen"] = ""
    return d


def _validar_items(raw_items: Any) -> list[dict[str, Any]]:
    if not isinstance(raw_items, list) or len(raw_items) == 0:
        raise ValueError("Agregá al menos un ítem al pedido.")
    out: list[dict[str, Any]] = []
    for i, it in enumerate(raw_items):
        if not isinstance(it, dict):
            raise ValueError(f"Ítem {i + 1} inválido.")
        descripcion = str(it.get("descripcion") or "").strip()
        if not descripcion:
            raise ValueError(f"Ítem {i + 1}: la descripción es obligatoria.")
        try:
            cantidad = float(it.get("cantidad") or 0)
        except (TypeError, ValueError) as e:
            raise ValueError(f"Ítem {i + 1}: cantidad inválida.") from e
        if cantidad <= 0:
            raise ValueError(f"Ítem {i + 1}: la cantidad debe ser mayor a 0.")
        area = str(it.get("area") or "").strip()
        if not area:
            raise ValueError(f"Ítem {i + 1}: área / máquina es obligatoria.")
        out.append(
            {
                "orden": i,
                "codigo": str(it.get("codigo") or "").strip(),
                "descripcion": descripcion,
                "cantidad": cantidad,
                "unidad": str(it.get("unidad") or "").strip(),
                "area": area,
                "imagen_path": str(it.get("imagen_path") or "").strip(),
            }
        )
    return out


# --------------------------------------------------------------------------- #
# Migración inicial desde SQLite (una sola vez)
# --------------------------------------------------------------------------- #
def _seed_desde_sqlite() -> None:
    """Si no existe el Excel pero hay un SQLite viejo con datos, lo importa."""
    if EXCEL_PATH.is_file() or not DB_PATH.is_file():
        return
    try:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        cols = {r[1] for r in conn.execute("PRAGMA table_info(solicitudes)").fetchall()}
        if not cols:
            conn.close()
            return
        filas = conn.execute("SELECT * FROM solicitudes ORDER BY id ASC").fetchall()
        solicitudes: list[dict[str, Any]] = []
        max_id = 0
        max_pedido = 0
        max_tr = 0
        for row in filas:
            d = dict(row)
            sid = int(d["id"])
            max_id = max(max_id, sid)
            items_rows = conn.execute(
                "SELECT * FROM solicitud_items WHERE solicitud_id = ? ORDER BY orden, id",
                (sid,),
            ).fetchall()
            items = [
                {
                    "orden": int(ir["orden"] or idx),
                    "codigo": str(ir["codigo"] or ""),
                    "descripcion": str(ir["descripcion"] or ""),
                    "cantidad": float(ir["cantidad"] or 0),
                    "unidad": str(ir["unidad"] or ""),
                    "area": str(ir["area"] or ""),
                    "imagen_path": str(ir["imagen_path"] or ""),
                }
                for idx, ir in enumerate(items_rows)
            ]
            sol = {"id": sid, "items": items, "proveedor": ""}
            for key in CABECERA_KEYS:
                if key == "proveedor":
                    continue
                sol[key] = str(d.get(key) or "")
            solicitudes.append(sol)
            np = str(d.get("n_pedido") or "")
            if np.upper().startswith("P-"):
                try:
                    max_pedido = max(max_pedido, int(np.split("-", 1)[1]))
                except (ValueError, IndexError):
                    pass
            ntr = str(d.get("n_tr") or "")
            if ntr.upper().startswith("TR-"):
                try:
                    max_tr = max(max_tr, int(ntr.split("-", 1)[1]))
                except (ValueError, IndexError):
                    pass
        conn.close()
        contadores = {"id_seq": max_id, "pedido": max_pedido, "tr": max_tr}
        with _FileLock(EXCEL_PATH):
            if not EXCEL_PATH.is_file():
                _escribir_workbook(solicitudes, contadores)
    except sqlite3.Error:
        return


def init_db() -> None:
    _asegurar_dirs()
    _seed_desde_sqlite()
    if not EXCEL_PATH.is_file():
        with _FileLock(EXCEL_PATH):
            if not EXCEL_PATH.is_file():
                _escribir_workbook([], {"id_seq": 0, "pedido": 0, "tr": 0})
    _load_all(force=True)


# --------------------------------------------------------------------------- #
# API pública (mismo contrato que la versión SQLite)
# --------------------------------------------------------------------------- #
def listar(
    *,
    tipo: str | None = None,
    estado: str | None = None,
    q: str | None = None,
    limite: int = 200,
) -> list[dict[str, Any]]:
    solicitudes, _ = _load_all()
    res = [_enrich(s) for s in solicitudes]

    if tipo == "tr":
        res = [s for s in res if s.get("tipo") == "tr"]
    elif tipo == "pedidos":
        res = [s for s in res if s.get("tipo") in ("normal", "urgente")]
    elif tipo:
        res = [s for s in res if s.get("tipo") == tipo]

    if estado:
        res = [s for s in res if s.get("estado") == estado]

    if q and q.strip():
        needle = q.strip().lower()

        def coincide(s: dict[str, Any]) -> bool:
            campos = [
                s.get("n_pedido", ""),
                s.get("n_tr", ""),
                s.get("solicitante", ""),
                s.get("cuenta_contable", ""),
                s.get("proveedor", ""),
                s.get("notas", ""),
            ]
            for it in s.get("items", []):
                campos.extend([it.get("codigo", ""), it.get("descripcion", ""), it.get("area", "")])
            return any(needle in str(v).lower() for v in campos)

        res = [s for s in res if coincide(s)]

    res.sort(key=lambda s: (str(s.get("creado_en") or ""), int(s.get("id") or 0)), reverse=True)
    return res[: max(1, int(limite))]


def resumen() -> dict[str, Any]:
    solicitudes, _ = _load_all()
    por_estado: dict[str, int] = {}
    urgentes = 0
    tr_abiertos = 0
    for s in solicitudes:
        est = str(s.get("estado") or "")
        por_estado[est] = por_estado.get(est, 0) + 1
        activo = est not in ("cumplido", "cancelado")
        if s.get("tipo") == "urgente" and activo:
            urgentes += 1
        if s.get("tipo") == "tr" and activo:
            tr_abiertos += 1
    return {
        "total": len(solicitudes),
        "por_estado": por_estado,
        "urgentes_activos": urgentes,
        "tr_activos": tr_abiertos,
    }


def obtener(solicitud_id: int) -> dict[str, Any] | None:
    solicitudes, _ = _load_all()
    for s in solicitudes:
        if int(s.get("id") or 0) == int(solicitud_id):
            return _enrich(s)
    return None


def proveedores_sugeridos() -> list[str]:
    solicitudes, _ = _load_all()
    vistos = {
        str(s.get("proveedor") or "").strip()
        for s in solicitudes
        if str(s.get("proveedor") or "").strip()
    }
    cats = cargar_catalogos()
    for p in cats.get("proveedores", []) or []:
        if str(p).strip():
            vistos.add(str(p).strip())
    return sorted(vistos, key=str.lower)


def crear(data: dict[str, Any], *, rol: str = "") -> dict[str, Any]:
    cats = cargar_catalogos()
    tipos_ok = {t["id"] for t in cats.get("tipos", [])}
    estados_ok = {e["id"] for e in cats.get("estados", [])}

    tipo = str(data.get("tipo") or "normal").strip().lower()
    if tipo not in tipos_ok:
        raise ValueError(f"Tipo inválido: {tipo}")

    cuenta = str(data.get("cuenta_contable") or "").strip()
    if not cuenta:
        raise ValueError("Cuenta contable es obligatoria.")

    solicitante = str(data.get("solicitante") or "").strip()
    if not solicitante:
        raise ValueError("Solicitante es obligatorio.")

    proveedor = str(data.get("proveedor") or "").strip()
    if tipo == "tr" and not proveedor:
        raise ValueError("En un TR el proveedor es obligatorio.")

    estado = str(data.get("estado") or "en_proceso").strip().lower()
    if estado not in estados_ok:
        estado = "en_proceso"

    items = _validar_items(data.get("items"))
    ahora = _ahora_iso()

    with _FileLock(EXCEL_PATH):
        solicitudes, contadores = _load_all(force=True)
        sid = _next(contadores, "id_seq")
        n_pedido = ""
        n_tr = ""
        if tipo in ("normal", "urgente"):
            n_pedido = f"P-{_next(contadores, 'pedido'):04d}"
        elif tipo == "tr":
            n_tr = f"TR-{_next(contadores, 'tr'):04d}"

        sol = {
            "id": sid,
            "n_pedido": n_pedido,
            "n_tr": n_tr,
            "tipo": tipo,
            "estado": estado,
            "cuenta_contable": cuenta,
            "solicitante": solicitante,
            "proveedor": proveedor,
            "remito_nro": str(data.get("remito_nro") or "").strip(),
            "remito_archivo": str(data.get("remito_archivo") or "").strip(),
            "presupuesto_nro": str(data.get("presupuesto_nro") or "").strip(),
            "presupuesto_archivo": str(data.get("presupuesto_archivo") or "").strip(),
            "notas": str(data.get("notas") or "").strip(),
            "creado_por": str(data.get("creado_por") or "").strip(),
            "creado_en": ahora,
            "actualizado_en": ahora,
            "items": items,
        }
        solicitudes.append(sol)
        _escribir_workbook(solicitudes, contadores)
    return _enrich(sol)


def actualizar(solicitud_id: int, data: dict[str, Any], *, rol: str = "") -> dict[str, Any] | None:
    cats = cargar_catalogos()
    estados_ok = {e["id"] for e in cats.get("estados", [])}
    tipos_ok = {t["id"] for t in cats.get("tipos", [])}
    puede_admin = rol in ("admin", "panol")

    with _FileLock(EXCEL_PATH):
        solicitudes, contadores = _load_all(force=True)
        sol = next((s for s in solicitudes if int(s.get("id") or 0) == int(solicitud_id)), None)
        if sol is None:
            return None

        cambios = False

        # Campos de texto que puede editar cualquier usuario con escritura,
        # incluidos remito y presupuesto (los puede cargar cualquiera).
        for key in (
            "cuenta_contable",
            "solicitante",
            "proveedor",
            "notas",
            "remito_nro",
            "presupuesto_nro",
            "remito_archivo",
            "presupuesto_archivo",
        ):
            if key in data:
                sol[key] = str(data[key] or "").strip()
                cambios = True

        if "tipo" in data:
            tipo = str(data["tipo"] or "").strip().lower()
            if tipo not in tipos_ok:
                raise ValueError(f"Tipo inválido: {tipo}")
            sol["tipo"] = tipo
            # Asignar correlativo si pasa a un tipo que aún no lo tiene.
            if tipo in ("normal", "urgente") and not str(sol.get("n_pedido") or "").strip():
                sol["n_pedido"] = f"P-{_next(contadores, 'pedido'):04d}"
            if tipo == "tr" and not str(sol.get("n_tr") or "").strip():
                sol["n_tr"] = f"TR-{_next(contadores, 'tr'):04d}"
            cambios = True

        if "estado" in data:
            estado = str(data["estado"] or "").strip().lower()
            if estado not in estados_ok:
                raise ValueError(f"Estado inválido: {estado}")
            sol["estado"] = estado
            cambios = True

        if "n_pedido" in data:
            if not puede_admin:
                raise ValueError("Sin permiso para editar Nº pedido.")
            sol["n_pedido"] = str(data["n_pedido"] or "").strip()
            cambios = True

        if str(sol.get("tipo")) == "tr" and not str(sol.get("proveedor") or "").strip():
            raise ValueError("En un TR el proveedor es obligatorio.")

        if "items" in data:
            sol["items"] = _validar_items(data.get("items"))
            cambios = True

        if cambios:
            sol["actualizado_en"] = _ahora_iso()
            _escribir_workbook(solicitudes, contadores)

    return _enrich(sol)


def eliminar(solicitud_id: int, *, rol: str = "") -> bool:
    if rol != "admin":
        raise ValueError("Solo un administrador puede borrar pedidos.")
    with _FileLock(EXCEL_PATH):
        solicitudes, contadores = _load_all(force=True)
        antes = len(solicitudes)
        solicitudes = [s for s in solicitudes if int(s.get("id") or 0) != int(solicitud_id)]
        if len(solicitudes) == antes:
            return False
        _escribir_workbook(solicitudes, contadores)
    return True


def ruta_upload(nombre: str) -> Path:
    UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
    return UPLOADS_DIR / nombre
