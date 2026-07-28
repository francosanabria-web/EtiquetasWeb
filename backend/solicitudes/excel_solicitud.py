# -*- coding: utf-8 -*-
"""Exporta una solicitud a un Excel prolijo (plantilla descargable)."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_AZUL = "1F4E78"
_AZUL_CLARO = "DDEBF7"
_GRIS = "F2F2F2"

_thin = Side(style="thin", color="BFBFBF")
_borde = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _num(id_: Any, sol: dict[str, Any]) -> str:
    if sol.get("n_tr"):
        return str(sol["n_tr"])
    if sol.get("n_pedido"):
        return str(sol["n_pedido"])
    return f"#{id_}"


def generar_excel(solicitud: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Solicitud"
    ws.sheet_view.showGridLines = False

    sid = solicitud.get("id")
    ws.merge_cells("A1:E1")
    t = ws["A1"]
    t.value = f"Solicitud de pedido  ·  {_num(sid, solicitud)}"
    t.font = Font(bold=True, size=16, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=_AZUL)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 28

    cabecera = [
        ("ID interno", f"#{sid}"),
        ("Nº pedido", solicitud.get("n_pedido")),
        ("Nº TR", solicitud.get("n_tr")),
        ("Tipo", solicitud.get("tipo")),
        ("Estado", solicitud.get("estado")),
        ("Cuenta contable", solicitud.get("cuenta_contable")),
        ("Solicitante", solicitud.get("solicitante")),
        ("Proveedor", solicitud.get("proveedor")),
        ("Nº remito", solicitud.get("remito_nro")),
        ("Nº presupuesto", solicitud.get("presupuesto_nro")),
        ("Creado", solicitud.get("creado_en")),
        ("Notas", solicitud.get("notas")),
    ]
    fila = 3
    for label, val in cabecera:
        c_lbl = ws.cell(row=fila, column=1, value=label)
        c_lbl.font = Font(bold=True, color=_AZUL)
        c_lbl.alignment = Alignment(vertical="center")
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=5)
        c_val = ws.cell(row=fila, column=2, value=str(val or "—"))
        c_val.alignment = Alignment(vertical="center", wrap_text=True)
        fila += 1

    fila += 1
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=5)
    h = ws.cell(row=fila, column=1, value="Ítems del pedido")
    h.font = Font(bold=True, size=12, color="FFFFFF")
    h.fill = PatternFill("solid", fgColor=_AZUL)
    h.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[fila].height = 22
    fila += 1

    headers = ["Código", "Descripción", "Cantidad", "Unidad", "Área / máquina"]
    for col, htxt in enumerate(headers, start=1):
        c = ws.cell(row=fila, column=col, value=htxt)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=_AZUL_CLARO)
        c.border = _borde
        c.alignment = Alignment(horizontal="center", vertical="center")
    fila += 1

    items = solicitud.get("items") or []
    for i, it in enumerate(items):
        vals = [
            it.get("codigo") or "—",
            it.get("descripcion") or "—",
            it.get("cantidad") or 0,
            it.get("unidad") or "—",
            it.get("area") or "—",
        ]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=fila, column=col, value=v)
            c.border = _borde
            c.alignment = Alignment(
                vertical="center",
                horizontal="center" if col in (1, 3, 4) else "left",
                wrap_text=(col == 2),
            )
            if i % 2 == 1:
                c.fill = PatternFill("solid", fgColor=_GRIS)
        fila += 1

    anchos = [16, 46, 10, 10, 26]
    for i, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    out = BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()
