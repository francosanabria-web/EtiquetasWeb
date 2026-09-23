# -*- coding: utf-8 -*-
"""IO Excel para maestro + historial de salidas (capa reemplazable por SQL)."""

from __future__ import annotations

import os
import shutil
import time
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import (
    ATENCIONES_COLUMNAS_EXPORT,
    COLUMNAS_MOVIMIENTO,
    HOJA_ARTICULOS,
    HOJA_CONFIG,
    HOJA_MOVIMIENTOS,
    MES_NOMBRES,
    OPERARIOS_DEFAULT,
    OPERARIOS_PROYECTOS_DEFAULT,
    SECTORES_DEFAULT,
    TIPOS_COMPROBANTE_DEFAULT,
    SALIDAS_ESCRIBIR_DIARIO,
    SALIDAS_DB_ENABLED,
    data_path,
    diario_path,
    historial_path,
    maestro_path,
    maestro_writable,
)

# Estética alineada a diarios de escritorio (salidas_DD-MM-AAAA.xlsx de planta).
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_ANCHOS_MOV = {
    "FECHA": 12.7,
    "MES": 12.0,
    "AÑO": 8.0,
    "CODIGO": 14.0,
    "DESCRIPCION": 45.7,
    "UBICACION": 12.7,
    "CANTIDAD": 12.0,
    "TIPO_COMPROBANTE": 18.7,
    "NUMERO_ORDEN": 14.7,
    "MAQUINA_SITIO": 27.7,
    "PRECIO_UNITARIO": 17.7,
    "MONTO_TOTAL_SALIDA": 20.7,
    "OPERARIO": 18.7,
    "SECTOR": 16.7,
}
_MONEY_COLS = frozenset({"PRECIO_UNITARIO", "MONTO_TOTAL_SALIDA"})
_INT_COLS = frozenset({"AÑO", "NUMERO_ORDEN"})

_LOCK_TIMEOUT_S = 12.0
_LOCK_STALE_S = 60.0


class FileLock:
    def __init__(self, target: Path) -> None:
        self.lock_path = target.with_suffix(target.suffix + ".lock")
        self._fd: int | None = None

    def __enter__(self) -> "FileLock":
        inicio = time.time()
        while True:
            try:
                self._fd = os.open(str(self.lock_path), os.O_CREAT | os.O_EXCL | os.O_RDWR)
                os.write(self._fd, str(os.getpid()).encode("ascii", "ignore"))
                return self
            except FileExistsError:
                try:
                    edad = time.time() - self.lock_path.stat().st_mtime
                    if edad > _LOCK_STALE_S:
                        self.lock_path.unlink(missing_ok=True)
                        continue
                except OSError:
                    pass
                if time.time() - inicio > _LOCK_TIMEOUT_S:
                    raise TimeoutError(
                        f"Archivo ocupado: {self.lock_path.name}. "
                        "Cerrá el Excel si está abierto e intentá de nuevo."
                    )
                time.sleep(0.15)

    def __exit__(self, *_) -> None:
        if self._fd is not None:
            try:
                os.close(self._fd)
            except OSError:
                pass
            self._fd = None
        try:
            self.lock_path.unlink(missing_ok=True)
        except OSError:
            pass


def _norm_header(c: Any) -> str:
    return str(c or "").strip().lower().replace(" ", "_").replace(".", "")


def fecha_sin_hora_str(fecha_val: Any) -> str:
    """Fecha escrita tipo 15/5/2026 (sin cero a la izquierda), como el escritorio."""
    if fecha_val is None:
        return ""
    try:
        if isinstance(fecha_val, pd.Timestamp):
            if pd.isna(fecha_val):
                return ""
            fecha_val = fecha_val.date()
        elif isinstance(fecha_val, datetime):
            fecha_val = fecha_val.date()
        if isinstance(fecha_val, date):
            return f"{fecha_val.day}/{fecha_val.month}/{fecha_val.year}"
        if isinstance(fecha_val, str):
            txt = fecha_val.strip()
            if not txt or txt.lower() in ("nan", "nat", "none"):
                return ""
            dt = pd.to_datetime(txt, dayfirst=True, errors="coerce")
            if pd.isna(dt):
                return txt
            return f"{dt.day}/{dt.month}/{dt.year}"
        if isinstance(fecha_val, (int, float)) and 20000 <= float(fecha_val) <= 120000:
            base = pd.Timestamp("1899-12-30")
            dt = base + pd.Timedelta(days=float(fecha_val))
            return f"{dt.day}/{dt.month}/{dt.year}"
    except Exception:
        pass
    return str(fecha_val).strip()


def parse_fecha(raw: str | date | datetime | None) -> date | None:
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    txt = str(raw).strip()
    if not txt:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(txt, fmt).date()
        except ValueError:
            continue
    dt = pd.to_datetime(txt, dayfirst=True, errors="coerce")
    if pd.isna(dt):
        return None
    return dt.date()


def nombre_mes(mes: int) -> str:
    return MES_NOMBRES.get(int(mes), "")


def normalizar_orden(val: Any) -> Any:
    s = str(val or "").strip()
    if not s:
        return ""
    s = s.replace(",", ".")
    try:
        f = float(s)
        if f == int(f):
            return int(f)
        return f
    except ValueError:
        return s


def col_stock(df: pd.DataFrame) -> str | None:
    for c in df.columns:
        n = _norm_header(c)
        if "stock" in n or n in ("act", "cant", "cantidad", "stkactual", "stockactual"):
            if "min" in n:
                continue
            return c
    return None


def col_precio(df: pd.DataFrame) -> str | None:
    for c in df.columns:
        n = _norm_header(c)
        if n in ("precio_unitario", "preciounitario", "costouni", "pu", "punit"):
            return c
        if "precio" in n:
            return c
    return None


def col_ubicacion(df: pd.DataFrame) -> str | None:
    for c in df.columns:
        n = _norm_header(c)
        if "ubic" in n:
            return c
    return None


def col_descripcion(df: pd.DataFrame) -> str | None:
    for c in df.columns:
        n = _norm_header(c)
        if n in ("descripcion", "descripción", "articulo", "artículo"):
            return c
    return None


def col_codigo(df: pd.DataFrame) -> str | None:
    for c in df.columns:
        if _norm_header(c) == "codigo":
            return c
    return None


def col_categoria(df: pd.DataFrame) -> str | None:
    for c in df.columns:
        n = _norm_header(c)
        if n in ("categoria", "categoría", "categoria_articulo"):
            return c
    return None


def asegurar_rutas_escritura() -> None:
    """Asegura carpeta de escritura + historial vacío.

    Nunca crea ni toca master_codes/master_salidas de producción del pañol.
    Solo genera un maestro DEMO si el destino del maestro es data_prueba y falta.
    """
    base = data_path()
    base.mkdir(parents=True, exist_ok=True)

    path = maestro_path()
    # Solo sembrar DEMO si el maestro faltante caería en data_prueba (modo local).
    if not path.is_file() and maestro_writable():
        path.parent.mkdir(parents=True, exist_ok=True)
        df_art = pd.DataFrame(
            [
                {
                    "codigo": "DEMO-001",
                    "descripcion": "TORNILLO HEX M8x20",
                    "ubicacion": "A-01",
                    "stock_actual": 100,
                    "precio_unitario": 15.5,
                    "categoria": "GENERAL",
                },
                {
                    "codigo": "DEMO-002",
                    "descripcion": "ARANDELA PLANA M8",
                    "ubicacion": "A-02",
                    "stock_actual": 250,
                    "precio_unitario": 2.0,
                    "categoria": "GENERAL",
                },
                {
                    "codigo": "DEMO-003",
                    "descripcion": "GRASA MULTIUSO 1KG",
                    "ubicacion": "B-10",
                    "stock_actual": 12,
                    "precio_unitario": 4500.0,
                    "categoria": "LUBRICANTES",
                },
            ]
        )
        n = max(
            7,
            len(TIPOS_COMPROBANTE_DEFAULT),
            len(OPERARIOS_DEFAULT),
        )
        sectores_cfg = [
            "MANTENIMIENTO",
            "MANTENIMIENTO",
            "EDILICIO",
            "PRODUCCION",
            "PROYECTOS",
            "PROYECTOS",
            "AUTOELEVADORES",
        ]
        operarios_cfg = ["AA", "AB", "AC", "AD", "MACCARONI", "VALENZUELA", "AE"]
        df_cfg = pd.DataFrame(
            {
                "sector": (sectores_cfg + [""] * n)[:n],
                "operario": (operarios_cfg + [""] * n)[:n],
                "comprobantes": (TIPOS_COMPROBANTE_DEFAULT + [""] * n)[:n],
                "operarios": (OPERARIOS_DEFAULT + [""] * n)[:n],
            }
        )
        with pd.ExcelWriter(path, engine="openpyxl") as w:
            df_art.to_excel(w, index=False, sheet_name=HOJA_ARTICULOS)
            df_cfg.to_excel(w, index=False, sheet_name=HOJA_CONFIG)

    hist = historial_path()
    if not hist.is_file():
        _escribir_excel_atomico(
            hist, pd.DataFrame(columns=COLUMNAS_MOVIMIENTO), HOJA_MOVIMIENTOS
        )


def asegurar_data_prueba() -> None:
    """Alias histórico — ahora asegura rutas de escritura (salidas_web o data_prueba)."""
    asegurar_rutas_escritura()


def leer_maestro() -> tuple[dict[str, pd.DataFrame], pd.DataFrame]:
    """Devuelve (hojas_por_nombre, config_df)."""
    asegurar_data_prueba()
    path = maestro_path()
    if not path.is_file():
        raise FileNotFoundError(f"No se encontró maestro: {path}")
    xls = pd.ExcelFile(path, engine="openpyxl")
    hojas: dict[str, pd.DataFrame] = {}
    config_df = pd.DataFrame()
    for name in xls.sheet_names:
        df = pd.read_excel(xls, sheet_name=name)
        key = str(name).strip()
        if _norm_header(key) == "config":
            config_df = df
        else:
            # Normalizar columna codigo a minúsculas internas para búsqueda
            c_cod = col_codigo(df)
            if c_cod and c_cod != "codigo":
                df = df.rename(columns={c_cod: "codigo"})
            if "codigo" in df.columns:
                df["codigo"] = df["codigo"].astype(str).str.strip().str.upper()
            hojas[key] = df
    if HOJA_ARTICULOS not in hojas and hojas:
        # Primera hoja con codigo
        for k, df in hojas.items():
            if "codigo" in df.columns:
                hojas[HOJA_ARTICULOS] = df
                break
    return hojas, config_df


def guardar_maestro(hojas: dict[str, pd.DataFrame], config_df: pd.DataFrame) -> None:
    """Persiste stock en el maestro SOLO si es writable (nunca producción)."""
    if not maestro_writable():
        return
    path = maestro_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with FileLock(path):
        try:
            with pd.ExcelWriter(tmp, engine="openpyxl") as w:
                for name, df in hojas.items():
                    df.to_excel(w, index=False, sheet_name=name[:31] or "Hoja")
                if config_df is not None and not config_df.empty:
                    config_df.to_excel(w, index=False, sheet_name=HOJA_CONFIG)
            if path.is_file():
                bak = path.with_suffix(path.suffix + ".bak")
                try:
                    shutil.copy2(path, bak)
                except OSError:
                    pass
            os.replace(tmp, path)
        except Exception:
            try:
                tmp.unlink(missing_ok=True)
            except OSError:
                pass
            raise


def leer_movimientos(path: Path | None = None) -> pd.DataFrame:
    p = path or historial_path()
    if not p.is_file():
        return pd.DataFrame(columns=COLUMNAS_MOVIMIENTO)
    df = pd.read_excel(p, engine="openpyxl")
    if "FECHA" in df.columns:
        df["FECHA"] = df["FECHA"].apply(fecha_sin_hora_str)
    return df


def _coerce_celda(col: str, val: Any) -> Any:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    if isinstance(val, pd.Timestamp):
        return fecha_sin_hora_str(val)
    if col == "FECHA":
        return fecha_sin_hora_str(val)
    if col in _MONEY_COLS or col == "CANTIDAD":
        try:
            return float(val) if str(val).strip() != "" else 0.0
        except (TypeError, ValueError):
            return val
    if col in _INT_COLS:
        try:
            f = float(val)
            return int(f) if f == int(f) else f
        except (TypeError, ValueError):
            return val
    return val


def _workbook_movimientos(df: pd.DataFrame, sheet_name: str) -> Workbook:
    """Workbook con cabecera/anchos/formatos como el Excel diario del escritorio."""
    cols = [c for c in COLUMNAS_MOVIMIENTO if c in df.columns] + [
        c for c in df.columns if c not in COLUMNAS_MOVIMIENTO
    ]
    if not cols:
        cols = list(COLUMNAS_MOVIMIENTO)

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or HOJA_MOVIMIENTOS)[:31]

    for col_idx, col in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _THIN
        ws.column_dimensions[get_column_letter(col_idx)].width = _ANCHOS_MOV.get(col, 14.0)

    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        for col_idx, (col, raw) in enumerate(zip(cols, row), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=_coerce_celda(col, raw))
            cell.border = _THIN
            if col in _MONEY_COLS:
                cell.number_format = "$ #,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col in _INT_COLS:
                cell.number_format = "0"
            elif col == "CANTIDAD":
                cell.alignment = Alignment(horizontal="right", vertical="center")

    last_row = max(1, len(df) + 1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last_row}"
    ws.row_dimensions[1].height = 18
    return wb


def _escribir_excel_atomico(path: Path, df: pd.DataFrame, sheet_name: str) -> None:
    """Escribe a .tmp y reemplaza — evita dejar el xlsx en 0 bytes si Drive/Excel interrumpe.

    Aplica formato profesional (cabecera azul, anchos, $ / enteros, freeze, filtro)
    alineado a los diarios de salidas del escritorio.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    try:
        wb = _workbook_movimientos(df, sheet_name)
        wb.save(tmp)
        if path.is_file():
            bak = path.with_suffix(path.suffix + ".bak")
            try:
                shutil.copy2(path, bak)
            except OSError:
                pass
        os.replace(tmp, path)
    except Exception:
        try:
            tmp.unlink(missing_ok=True)
        except OSError:
            pass
        raise


def append_movimientos(filas: list[dict[str, Any]], fechas: list[date]) -> list[str]:
    """Escribe en master_salidas + (opcional) Excel diario por fecha. Devuelve paths tocados.

    Comportamiento segun SALIDAS_ESCRIBIR_DIARIO:
      - false (default): solo escribe historial_path (master_salidas.xlsx) /
        salida_historial como fuente unica. El diario pasa a ser export
        on-demand via exportar_diario_excel() sin persistir archivo diario.
        Beneficios: menos fragmentacion, sin cientos de archivos, backup simple,
        menos locks .lock huérfanos en Drive, auditoria en una sola tabla.
      - true: escritura dual historial + salidas_DD-MM-AAAA.xlsx (legado, compat).
    """
    if not filas:
        return []
    asegurar_data_prueba()
    tocados: list[str] = []
    por_archivo: dict[Path, list[dict[str, Any]]] = {}
    hist = historial_path()
    por_archivo.setdefault(hist, []).extend(filas)
    if SALIDAS_ESCRIBIR_DIARIO:
        for fila, f in zip(filas, fechas):
            por_archivo.setdefault(diario_path(f), []).append(fila)

    for path, batch in por_archivo.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        with FileLock(path):
            if path.is_file() and path.stat().st_size >= 64:
                try:
                    df_p = pd.read_excel(path, engine="openpyxl")
                except Exception as e:
                    raise RuntimeError(
                        f"No se pudo leer {path.name} (archivo corrupto o vacío). "
                        f"Restaurá desde .bak o diarios antes de registrar salidas. Detalle: {e}"
                    ) from e
            else:
                df_p = pd.DataFrame(columns=COLUMNAS_MOVIMIENTO)
            if "FECHA" in df_p.columns:
                df_p["FECHA"] = df_p["FECHA"].apply(fecha_sin_hora_str)
            df_p = pd.concat([df_p, pd.DataFrame(batch)], ignore_index=True)
            if "FECHA" in df_p.columns:
                df_p["FECHA"] = df_p["FECHA"].apply(fecha_sin_hora_str)
            # Orden de columnas preferido
            cols = [c for c in COLUMNAS_MOVIMIENTO if c in df_p.columns] + [
                c for c in df_p.columns if c not in COLUMNAS_MOVIMIENTO
            ]
            df_p = df_p[cols]
            _escribir_excel_atomico(path, df_p, HOJA_MOVIMIENTOS)
        tocados.append(str(path))
    return tocados


def exportar_diario_excel(fecha: date | str, *, df_or_rows: pd.DataFrame | list[dict[str, Any]] | None = None) -> Workbook:
    """Genera Workbook diario on-demand sin persistir archivo.

    Fuente prioridad:
      1) df_or_rows si se pasa
      2) Si SALIDAS_DB_ENABLED -> salida_historial via store_sql.historial_listar(fecha)
      3) Sino lee master_salidas / historial_path (Excel) filtrando por fecha.

    Es la vista exportable que reemplaza los archivos
    salidas_DD-MM-AAAA.xlsx fragmentados. Usado por Reportes y por
    GET /api/salidas/diario/export?fecha=YYYY-MM-DD
    """
    d = parse_fecha(fecha) if isinstance(fecha, str) else fecha
    if d is None:
        raise ValueError("fecha invalida para export diario")
    if df_or_rows is None:
        # Intentar DB si flag activo
        if SALIDAS_DB_ENABLED:
            try:
                from store_sql import historial_listar

                rows = historial_listar(desde=d, hasta=d, limite=10000)
                if rows:
                    # Mapear rows DB (snake_case) a COLUMNAS_MOVIMIENTO (UPPER)
                    mapped = []
                    for r in rows:
                        mapped.append(
                            {
                                "FECHA": fecha_sin_hora_str(r.get("fecha")),
                                "MES": r.get("mes") or nombre_mes(d.month),
                                "AÑO": r.get("anio") or d.year,
                                "CODIGO": r.get("codigo"),
                                "DESCRIPCION": r.get("descripcion"),
                                "UBICACION": r.get("ubicacion"),
                                "CANTIDAD": r.get("cantidad"),
                                "TIPO_COMPROBANTE": r.get("tipo_comprobante"),
                                "NUMERO_ORDEN": r.get("numero_orden"),
                                "MAQUINA_SITIO": r.get("maquina_sitio"),
                                "PRECIO_UNITARIO": r.get("precio_unitario"),
                                "MONTO_TOTAL_SALIDA": r.get("monto_total"),
                                "OPERARIO": r.get("operario_nombre"),
                                "SECTOR": r.get("sector_nombre"),
                            }
                        )
                    df = pd.DataFrame(mapped)
                    if df.empty:
                        df = pd.DataFrame(columns=COLUMNAS_MOVIMIENTO)
                    return _workbook_movimientos(df, HOJA_MOVIMIENTOS)
            except Exception:
                pass
        df = leer_movimientos()
        if not df.empty and "FECHA" in df.columns:
            fechas = df["FECHA"].apply(parse_fecha)
            df = df[fechas == d].reset_index(drop=True)
        else:
            df = pd.DataFrame(columns=COLUMNAS_MOVIMIENTO)
    elif isinstance(df_or_rows, pd.DataFrame):
        df = df_or_rows.copy()
    else:
        df = pd.DataFrame(df_or_rows)
        if df.empty:
            df = pd.DataFrame(columns=COLUMNAS_MOVIMIENTO)
    return _workbook_movimientos(df, HOJA_MOVIMIENTOS)


def generar_diario_bytes(fecha: date | str) -> tuple[bytes, str]:
    """Retorna (bytes, filename) para diario export, soporta DB o Excel.

    Usa exportar_diario_excel internamente.
    """
    import io

    d = parse_fecha(fecha) if isinstance(fecha, str) else fecha
    if d is None:
        raise ValueError("fecha invalida")
    wb = exportar_diario_excel(d)
    buf = io.BytesIO()
    wb.save(buf)
    fname = f"salidas_{d.strftime('%d-%m-%Y')}.xlsx"
    return buf.getvalue(), fname


# ---------- Remito / Comprobante de salida ----------

_REMITO_FILL_HEADER = PatternFill("solid", fgColor="1F4E78")
_REMITO_FILL_TOTAL = PatternFill("solid", fgColor="D9E1F2")
_REMITO_FILL_FIRMA = PatternFill("solid", fgColor="F2F2F2")


def _workbook_remito(items: list[dict[str, Any]], cabecera: dict[str, Any]) -> Workbook:
    """Workbook Remito / Comprobante de salida con header, tabla items, total y firmas."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Remito"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True

    # Anchuras
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 15
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Header empresa
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "PAÑOL — COMPROBANTE DE SALIDA / REMITO"
    c.font = Font(bold=True, size=13, color="1F4E78")
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.fill = PatternFill("solid", fgColor="EAF0F7")
    ws.row_dimensions[1].height = 22

    # Sub-header orden / fecha
    fecha_s = fecha_sin_hora_str(cabecera.get("fecha")) or str(cabecera.get("fecha") or "")
    ws.merge_cells("A2:F2")
    ws["A2"].value = f"Nº ORDEN: {cabecera.get('numero_orden') or cabecera.get('orden') or '—'}    •    FECHA: {fecha_s or '—'}    •    COMPROBANTE: {cabecera.get('tipo_comprobante') or '—'}"
    ws["A2"].font = Font(bold=True, size=10)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A2"].fill = PatternFill("solid", fgColor="F2F2F2")
    ws.row_dimensions[2].height = 16

    # Datos cabecera en dos filas
    row = 3
    headers = ["Sector:", "Operario:", "Máquina/Sitio:"]
    values = [
        str(cabecera.get("sector") or cabecera.get("sector_nombre") or "—").upper(),
        str(cabecera.get("operario") or cabecera.get("operario_nombre") or "—").upper(),
        str(cabecera.get("maquina") or cabecera.get("maquina_sitio") or "—").upper(),
    ]
    # Fila 3: labels
    for idx, h in enumerate(headers, start=1):
        col = get_column_letter(idx * 2 - 1)
        # merge each pair A-B, C-D, E-F
        start_col = (idx - 1) * 2 + 1
        end_col = start_col + 1
        ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
        cell = ws.cell(row=row, column=start_col, value=h)
        cell.font = Font(bold=True, size=8, color="5A5A5A")
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 12
    row += 1
    # Fila 4: valores
    for idx, v in enumerate(values, start=1):
        start_col = (idx - 1) * 2 + 1
        end_col = start_col + 1
        ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
        cell = ws.cell(row=row, column=start_col, value=v)
        cell.font = Font(bold=True, size=10)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = Border(bottom=Side(style="thin", color="BFBFBF"))
    ws.row_dimensions[row].height = 15
    row += 1
    # Espacio
    row += 1  # fila 6 vacia

    # Tabla items header
    headers_items = ["CÓDIGO", "DESCRIPCIÓN", "UBICACIÓN", "CANT.", "P. UNIT.", "MONTO"]
    item_header_row = row
    for col_idx, h in enumerate(headers_items, start=1):
        cell = ws.cell(row=row, column=col_idx, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN
    ws.row_dimensions[row].height = 16
    row += 1

    total_monto = 0.0
    for it in items or []:
        codigo = str(it.get("CODIGO") or it.get("codigo") or "").upper()
        desc = str(it.get("DESCRIPCION") or it.get("descripcion") or "").upper()
        ubic = str(it.get("UBICACION") or it.get("ubicacion") or "").upper()
        try:
            cant = float(it.get("CANTIDAD") if "CANTIDAD" in it else it.get("cantidad") or 0)
        except Exception:
            cant = 0
        try:
            pu = float(it.get("PRECIO_UNITARIO") if "PRECIO_UNITARIO" in it else it.get("precio_unitario") or 0)
        except Exception:
            pu = 0
        try:
            monto = float(it.get("MONTO_TOTAL_SALIDA") if "MONTO_TOTAL_SALIDA" in it else it.get("monto_total") or (abs(cant) * abs(pu)))
        except Exception:
            monto = abs(cant) * abs(pu)
        total_monto += float(monto or 0)
        row_vals = [codigo, desc, ubic, cant, pu, monto]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.border = _THIN
            cell.font = Font(size=9)
            if col_idx in (4, 5, 6):
                cell.number_format = "#,##0.00" if col_idx == 4 else "$ #,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if col_idx == 1:
                cell.font = Font(size=9, bold=True)
        ws.row_dimensions[row].height = 14
        row += 1

    if not items:
        # fila vacia placeholder
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = ws.cell(row=row, column=1, value="Sin ítems")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = Font(italic=True, color="888888", size=9)
        cell.border = _THIN
        ws.row_dimensions[row].height = 14
        row += 1

    # Total
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    cell = ws.cell(row=row, column=1, value="TOTAL")
    cell.font = Font(bold=True, size=10, color="1F4E78")
    cell.alignment = Alignment(horizontal="right", vertical="center")
    cell.fill = _REMITO_FILL_TOTAL
    cell.border = _THIN
    for c in range(2, 6):
        ws.cell(row=row, column=c).fill = _REMITO_FILL_TOTAL
        ws.cell(row=row, column=c).border = _THIN
    cell_total = ws.cell(row=row, column=6, value=total_monto)
    cell_total.font = Font(bold=True, size=10, color="1F4E78")
    cell_total.number_format = "$ #,##0.00"
    cell_total.alignment = Alignment(horizontal="right", vertical="center")
    cell_total.fill = _REMITO_FILL_TOTAL
    cell_total.border = _THIN
    ws.row_dimensions[row].height = 16
    row += 2

    # Firmas
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
    c1 = ws.cell(row=row, column=1, value="Retiró (firma y aclaración)")
    c2 = ws.cell(row=row, column=4, value="Entregó — Pañol (firma)")
    for c in (c1, c2):
        c.font = Font(bold=True, size=8, color="5A5A5A")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.fill = _REMITO_FILL_FIRMA
        c.border = _THIN
    ws.row_dimensions[row].height = 14
    row += 1
    # lineas firma
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 2, end_column=3)
    ws.merge_cells(start_row=row, start_column=4, end_row=row + 2, end_column=6)
    for r in range(row, row + 3):
        for col in (1, 4):
            cc = ws.cell(row=r, column=col)
            cc.border = _THIN
            if r == row + 2:
                cc.alignment = Alignment(horizontal="center", vertical="bottom")
                cc.value = " " if r != row + 2 else None
        # fill merged area
    # height for signature area
    for r in range(row, row + 3):
        ws.row_dimensions[r].height = 18
    row += 3
    # footer
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    cell = ws.cell(row=row, column=1, value=f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}  •  Sistema Pañol — Salidas")
    cell.font = Font(italic=True, size=7, color="888888")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 10

    # Print area
    ws.print_area = f"A1:F{row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4

    return wb


def exportar_remito_excel(items: list[dict[str, Any]], cabecera: dict[str, Any]) -> Workbook:
    """Wrapper publico para remito."""
    return _workbook_remito(items, cabecera)


def generar_remito_bytes(orden: str, fecha: date | str | None = None, *, items: list[dict[str, Any]] | None = None, cabecera: dict[str, Any] | None = None) -> tuple[bytes, str]:
    """Genera remito bytes para una orden. Si no se pasan items/cabecera, los busca en DB/Excel.

    Retorna (bytes, filename).
    """
    import io

    orden_s = str(orden or "").strip()
    if not orden_s:
        raise ValueError("orden es obligatoria")
    d = parse_fecha(fecha) if isinstance(fecha, str) else fecha if isinstance(fecha, date) else None

    # Si no vienen items, buscar en DB o Excel
    if items is None:
        if SALIDAS_DB_ENABLED:
            try:
                from store_sql import historial_por_orden_fecha, historial_listar
                if d:
                    rows = historial_por_orden_fecha(orden_s, d)
                else:
                    rows = historial_por_orden_fecha(orden_s)
                if not rows:
                    # fallback buscar por q orden
                    rows = historial_listar(numero_orden=orden_s, limite=100)
                    if d:
                        rows = [r for r in rows if str(r.get("fecha")) == str(d)]
                # Mapear a COLUMNAS
                items = []
                for r in rows:
                    items.append(
                        {
                            "CODIGO": r.get("codigo"),
                            "DESCRIPCION": r.get("descripcion"),
                            "UBICACION": r.get("ubicacion"),
                            "CANTIDAD": r.get("cantidad"),
                            "PRECIO_UNITARIO": r.get("precio_unitario"),
                            "MONTO_TOTAL_SALIDA": r.get("monto_total"),
                            "FECHA": r.get("fecha"),
                            "TIPO_COMPROBANTE": r.get("tipo_comprobante"),
                            "NUMERO_ORDEN": r.get("numero_orden"),
                            "MAQUINA_SITIO": r.get("maquina_sitio"),
                            "OPERARIO": r.get("operario_nombre"),
                            "SECTOR": r.get("sector_nombre"),
                        }
                    )
                # Cabecera derivada del primer row si no viene
                if cabecera is None and rows:
                    r0 = rows[0]
                    cabecera = {
                        "fecha": r0.get("fecha"),
                        "numero_orden": r0.get("numero_orden"),
                        "tipo_comprobante": r0.get("tipo_comprobante"),
                        "sector": r0.get("sector_nombre"),
                        "operario": r0.get("operario_nombre"),
                        "maquina": r0.get("maquina_sitio"),
                    }
            except Exception:
                items = []
        if items is None or (not items and not SALIDAS_DB_ENABLED):
            # Excel fallback
            df = leer_movimientos()
            if not df.empty:
                # filtrar por orden
                if "NUMERO_ORDEN" in df.columns:
                    df_f = df[df["NUMERO_ORDEN"].astype(str).str.strip() == orden_s]
                    if d is not None and "FECHA" in df_f.columns:
                        df_f = df_f[df_f["FECHA"].apply(parse_fecha) == d]
                    if not df_f.empty:
                        items = df_f.to_dict(orient="records")
                        if cabecera is None:
                            row0 = df_f.iloc[0]
                            cabecera = {
                                "fecha": row0.get("FECHA"),
                                "numero_orden": row0.get("NUMERO_ORDEN"),
                                "tipo_comprobante": row0.get("TIPO_COMPROBANTE"),
                                "sector": row0.get("SECTOR"),
                                "operario": row0.get("OPERARIO"),
                                "maquina": row0.get("MAQUINA_SITIO"),
                            }
                    else:
                        items = []
                else:
                    items = []
            else:
                items = []
    if cabecera is None:
        cabecera = {"numero_orden": orden_s, "fecha": d or date.today()}
    # Asegurar fecha string para header
    if not cabecera.get("fecha") and d:
        cabecera["fecha"] = d
    wb = _workbook_remito(items or [], cabecera)
    buf = io.BytesIO()
    wb.save(buf)
    fecha_part = d.strftime("%Y-%m-%d") if isinstance(d, date) else str(cabecera.get("fecha") or "s-f")
    safe_orden = "".join(c if c.isalnum() else "_" for c in orden_s)[:20] or "remito"
    fname = f"remito_{safe_orden}_{fecha_part}.xlsx"
    return buf.getvalue(), fname


# ---------- Atenciones ventanilla - export helpers (v4 simplificado) ----------

_ANCHOS_ATENCION = {
    "FECHA": 13.0,
    "CON_RETIRO": 18.0,
    "OBSERVACIONES": 48.0,
    "ATENDIDO_POR": 16.0,
}


def _atencion_row_to_export(row: dict[str, Any]) -> dict[str, Any]:
    """Normaliza fila DB salida_atenciones -> columnas export v4 simplificadas."""
    fecha_v = row.get("fecha")
    if isinstance(fecha_v, (date, datetime)):
        fecha_s = f"{fecha_v.day}/{fecha_v.month}/{fecha_v.year}"
    else:
        fecha_s = fecha_sin_hora_str(fecha_v) if fecha_v else ""
    con = row.get("con_retiro")
    # Semantica v4: SI = con retiro FUERA DE SISTEMA, NO = sin stock
    if con == 1 or con is True or str(con).strip() in ("1", "true", "True", "SI"):
        con_s = "SI — fuera sistema"
    elif con == 0 or con is False or str(con).strip() in ("0", "false", "False", "NO"):
        con_s = "NO — sin stock"
    else:
        con_s = "NO — sin stock" if not con else "SI — fuera sistema"
    # Observaciones 500 + compat legacy
    obs = row.get("observaciones")
    if obs is None:
        # compat legacy: construir desde motivo/persona si existen
        legacy_parts = []
        if row.get("motivo_sin_retiro"):
            legacy_parts.append(str(row.get("motivo_sin_retiro")))
        if row.get("persona_solicitante"):
            legacy_parts.append(f"({row.get('persona_solicitante')})")
        obs = " ".join(legacy_parts) if legacy_parts else ""
    else:
        obs = str(obs).strip()
    return {
        "FECHA": fecha_s,
        "CON_RETIRO": con_s,
        "OBSERVACIONES": obs[:500],
        "ATENDIDO_POR": str(row.get("atendido_por") or row.get("creado_por") or "").strip().upper(),
    }


def workbook_atenciones(rows: list[dict[str, Any]], sheet_name: str = "Atenciones") -> Workbook:
    """Workbook con formato profesional para export atenciones."""
    cols = list(ATENCIONES_COLUMNAS_EXPORT)
    # Normalizar rows a columnas export
    norm_rows = [_atencion_row_to_export(r) for r in (rows or [])]
    df = pd.DataFrame(norm_rows, columns=cols) if norm_rows else pd.DataFrame(columns=cols)
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Atenciones")[:31]
    for col_idx, col in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _THIN
        ws.column_dimensions[get_column_letter(col_idx)].width = _ANCHOS_ATENCION.get(col, 14.0)
    for row_idx, rec in enumerate(norm_rows, start=2):
        for col_idx, col in enumerate(cols, start=1):
            raw = rec.get(col, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=raw)
            cell.border = _THIN
            if col in ("CANTIDAD_ITEMS_SOLICITADOS",):
                cell.alignment = Alignment(horizontal="center", vertical="center")
                try:
                    if raw != "":
                        cell.value = int(raw)
                        cell.number_format = "0"
                except (TypeError, ValueError):
                    pass
            elif col == "CON_RETIRO":
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if raw == "SI":
                    cell.fill = PatternFill("solid", fgColor="E8F5E9")
                elif raw == "NO":
                    cell.fill = PatternFill("solid", fgColor="FFEBEE")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
    last_row = max(1, len(norm_rows) + 1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last_row}"
    ws.row_dimensions[1].height = 18
    return wb


def atenciones_to_csv_bytes(rows: list[dict[str, Any]]) -> bytes:
    """CSV UTF-8 con BOM para Excel - columnas estables."""
    import csv
    import io

    cols = list(ATENCIONES_COLUMNAS_EXPORT)
    norm_rows = [_atencion_row_to_export(r) for r in (rows or [])]
    out = io.StringIO()
    # BOM para que Excel detecte UTF-8
    out.write("\ufeff")
    w = csv.DictWriter(out, fieldnames=cols, extrasaction="ignore", lineterminator="\r\n")
    w.writeheader()
    for rec in norm_rows:
        w.writerow(rec)
    return out.getvalue().encode("utf-8")


def catalogos_desde_config(config_df: pd.DataFrame) -> dict[str, Any]:
    """Lee hoja `config` de master_codes (comprobantes / sector / operario).

    Misma fuente que el escritorio: columnas `comprobantes`, `operarios`,
    y pares `sector`+`operario`. Solo lectura — no modifica el Excel.
    """
    tipos = list(TIPOS_COMPROBANTE_DEFAULT)
    sectores = list(SECTORES_DEFAULT)
    operarios = list(OPERARIOS_DEFAULT)
    operarios_proyectos = list(OPERARIOS_PROYECTOS_DEFAULT)
    sector_map: dict[str, list[str]] = {}

    if config_df is not None and not config_df.empty:
        cols = {_norm_header(c): c for c in config_df.columns}
        if "comprobantes" in cols:
            tipos = [
                str(x).strip().upper()
                for x in config_df[cols["comprobantes"]].dropna().tolist()
                if str(x).strip() and str(x).strip().lower() not in ("nan", "none")
            ] or tipos
        if "operarios" in cols:
            operarios = [
                str(x).strip().upper()
                for x in config_df[cols["operarios"]].dropna().tolist()
                if str(x).strip() and str(x).strip().lower() not in ("nan", "none")
            ] or operarios
        # Pares sector → operario (columna singular; no mezclar con lista `operarios`)
        col_sec = cols.get("sector")
        col_ope = cols.get("operario")
        if col_sec and col_ope:
            for _, row in config_df.iterrows():
                s = str(row.get(col_sec, "")).strip().upper()
                o = str(row.get(col_ope, "")).strip().upper()
                if not s or s in ("NAN", "NONE"):
                    continue
                sector_map.setdefault(s, [])
                if o and o not in ("NAN", "NONE") and o not in sector_map[s]:
                    sector_map[s].append(o)
            if sector_map:
                sectores = sorted(sector_map.keys())

    for s in SECTORES_DEFAULT:
        sector_map.setdefault(s, [])
    if "PROYECTOS" in sector_map and not sector_map["PROYECTOS"]:
        sector_map["PROYECTOS"] = list(operarios_proyectos)
    # Operarios de proyectos desde el mapa real si existen
    if sector_map.get("PROYECTOS"):
        operarios_proyectos = list(sector_map["PROYECTOS"])

    return {
        "tipos_comprobante": tipos,
        "sectores": sectores,
        "operarios": operarios,
        "operarios_proyectos": operarios_proyectos,
        "sector_operarios": sector_map,
    }
