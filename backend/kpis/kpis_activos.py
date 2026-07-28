# -*- coding: utf-8 -*-
"""KPIs de activos fuera de planta."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

import pandas as pd

from data_loader import DataStore

# Mismos umbrales/colores que el mail de seguimiento y la web.
_DIAS_AVISO = 21
_DIAS_ALERTA = 30
_VERDE = "C6EFCE"
_AMARILLO = "FFEB9C"
_ROJO = "FFC7CE"

_COLUMNAS_EXPORT = [
    ("equipo", "Equipo / repuesto"),
    ("codigo", "Código"),
    ("remito", "Nº remito"),
    ("n_pedido", "Nº pedido"),
    ("n_oc", "Nº OC"),
    ("sector", "Sector"),
    ("dias_fuera", "Días fuera"),
    ("proveedor", "Proveedor"),
    ("estado", "Estado"),
]


def _color_dias(dias: int) -> str:
    if dias > _DIAS_ALERTA:
        return _ROJO
    if dias > _DIAS_AVISO:
        return _AMARILLO
    return _VERDE


def _norm_doc(val: object) -> str:
    s = str(val or "").strip()
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s if s and s.lower() not in ("nan", "none") else ""


def _norm_text(val: object) -> str:
    s = str(val or "").strip()
    return "" if s.lower() in ("nan", "none", "nat") else s


def activos_resumen(store: DataStore) -> dict[str, Any]:
    df = store.activos
    if df.empty:
        return {
            "fuera_de_planta": 0,
            "dias_promedio_fuera": 0.0,
            "por_sector": [],
            "lista": [],
            "ultima_actualizacion": store.timestamp_iso(),
        }

    pend = df[df["_pendiente"]].copy()
    dias = pd.to_numeric(pend["DIAS_FUERA"], errors="coerce").fillna(0)
    dias_prom = round(float(dias.mean()), 1) if len(pend) else 0.0

    por_sector: list[dict[str, Any]] = []
    if "SECTOR" in pend.columns and not pend.empty:
        grp = pend.groupby("SECTOR", dropna=False).size().reset_index(name="cantidad")
        for _, r in grp.iterrows():
            por_sector.append({"sector": str(r["SECTOR"]), "cantidad": int(r["cantidad"])})

    lista = []
    for _, row in pend.sort_values("DIAS_FUERA", ascending=False).iterrows():
        lista.append(
            {
                "codigo": _norm_doc(row.get("CODIGO", "")),
                "equipo": _norm_text(row.get("EQUIPO_REPUESTO", "")),
                "sector": str(row.get("SECTOR", "")),
                "dias_fuera": int(float(row.get("DIAS_FUERA", 0) or 0)),
                "estado": str(row.get("ESTADO", "")),
                "proveedor": str(row.get("PROVEEDOR", "")),
                "remito": _norm_doc(row.get("NUMERO_REMITO", "")),
                "n_pedido": _norm_doc(row.get("NUMERO_PEDIDO", "")),
                "n_oc": _norm_doc(row.get("NUMERO_OC", "")),
            }
        )

    return {
        "fuera_de_planta": int(len(pend)),
        "dias_promedio_fuera": dias_prom,
        "por_sector": por_sector,
        "lista": lista,
        "ultima_actualizacion": store.timestamp_iso(),
    }


def _estado_legible(s: str) -> str:
    return str(s or "").replace("_", " ").strip() or "—"


def activos_excel(store: DataStore) -> bytes:
    """Export prolijo a Excel con filas coloreadas por días fuera."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    data = activos_resumen(store)
    lista = data.get("lista", [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Fuera de planta"
    ws.sheet_view.showGridLines = False

    n_cols = len(_COLUMNAS_EXPORT)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    t = ws.cell(row=1, column=1, value="Activos fuera de planta")
    t.font = Font(bold=True, size=15, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor="A32D2D")
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    sub = ws.cell(
        row=2,
        column=1,
        value=(
            f"{data.get('fuera_de_planta', 0)} equipos fuera · "
            f"promedio {data.get('dias_promedio_fuera', 0)} días · "
            f"generado {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ),
    )
    sub.font = Font(size=10, italic=True, color="555555")
    sub.alignment = Alignment(horizontal="left", indent=1)

    thin = Side(style="thin", color="BFBFBF")
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)

    fila = 4
    for col, (_key, titulo) in enumerate(_COLUMNAS_EXPORT, start=1):
        c = ws.cell(row=fila, column=col, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="475569")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = borde
    ws.freeze_panes = f"A{fila + 1}"

    for it in lista:
        fila += 1
        dias = int(it.get("dias_fuera", 0) or 0)
        color = _color_dias(dias)
        for col, (key, _titulo) in enumerate(_COLUMNAS_EXPORT, start=1):
            if key == "estado":
                val: Any = _estado_legible(it.get("estado", ""))
            elif key == "dias_fuera":
                val = dias
            else:
                val = it.get(key) or "—"
            c = ws.cell(row=fila, column=col, value=val)
            c.fill = PatternFill("solid", fgColor=color)
            c.border = borde
            c.alignment = Alignment(
                horizontal="center" if key in ("dias_fuera", "sector") else "left",
                vertical="center",
            )

    anchos = [30, 14, 14, 14, 12, 16, 11, 22, 18]
    for i, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    out = BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def activos_pdf(store: DataStore) -> bytes:
    """Export prolijo a PDF (apaisado) con la misma información y colores."""
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    data = activos_resumen(store)
    lista = data.get("lista", [])

    def s(v: Any, n: int | None = None) -> str:
        txt = str(v if v not in (None, "") else "-")
        txt = txt.encode("latin-1", "replace").decode("latin-1")
        if n and len(txt) > n:
            return txt[: n - 1] + "."
        return txt

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(0, 9, "Activos fuera de planta", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(
        0,
        6,
        s(
            f"{data.get('fuera_de_planta', 0)} equipos fuera  -  "
            f"promedio {data.get('dias_promedio_fuera', 0)} dias  -  "
            f"generado {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ),
        new_x=XPos.LMARGIN,
        new_y=YPos.NEXT,
    )
    pdf.ln(2)

    col_w = [58, 24, 24, 24, 20, 30, 20, 40, 13]
    titulos = [t for _k, t in _COLUMNAS_EXPORT]
    # Reordenar: Días al final en la web va antes de proveedor; acá mantenemos orden de columnas.
    pdf.set_font("Helvetica", "B", 8)
    pdf.set_fill_color(71, 85, 105)
    pdf.set_text_color(255, 255, 255)
    for w, tit in zip(col_w, titulos):
        pdf.cell(w, 7, s(tit, 26), border=1, fill=True, align="C", new_x=XPos.RIGHT, new_y=YPos.TOP)
    pdf.ln()

    pdf.set_font("Helvetica", "", 8)
    pdf.set_text_color(20, 20, 20)
    hex_rgb = {
        _VERDE: (198, 239, 206),
        _AMARILLO: (255, 235, 156),
        _ROJO: (255, 199, 206),
    }
    for it in lista:
        dias = int(it.get("dias_fuera", 0) or 0)
        r, g, b = hex_rgb[_color_dias(dias)]
        pdf.set_fill_color(r, g, b)
        celdas = [
            s(it.get("equipo"), 34),
            s(it.get("codigo"), 14),
            s(it.get("remito"), 14),
            s(it.get("n_pedido"), 14),
            s(it.get("n_oc"), 11),
            s(it.get("sector"), 16),
            str(dias),
            s(it.get("proveedor"), 24),
            s(_estado_legible(it.get("estado", "")), 8),
        ]
        for w, v in zip(col_w, celdas):
            pdf.cell(w, 6, v, border=1, fill=True, new_x=XPos.RIGHT, new_y=YPos.TOP)
        pdf.ln()

    out = BytesIO()
    pdf.output(out)
    return out.getvalue()
