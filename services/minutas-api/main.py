# -*- coding: utf-8 -*-
"""API REST del módulo Minutas de reunión semanal."""

from __future__ import annotations

import asyncio
import io
import os
import smtplib
from pathlib import Path

from fastapi import FastAPI, File, HTTPException, UploadFile, status, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from contextlib import asynccontextmanager

import repo
import repo_import
import openpyxl
from excel_parser import parsear_excel, exportar_excel, FilaPedido, _ref_pedido
from mail_builder import construir_mail_minuta_excel
from email_service import (
    EmailNoConfigurado,
    construir_asunto,
    construir_cuerpo_html,
    construir_cuerpo_texto,
    destinatarios_default_o,
    enviar_minuta,
)
from models import (
    Actualizacion,
    ActualizacionCreate,
    EnviarMinutaRequest,
    EnviarMinutaResponse,
    EntregaParcialCreate,
    PreviewEmailResponse,
    SesionCreate,
    SesionDetalle,
    SesionResumen,
    SesionUpdate,
    Solicitud,
    SolicitudCreate,
    SolicitudUpdate,
    Tema,
    TemaCreate,
    TemaUpdate,
)


def _cargar_env_local() -> None:
    """Lee services/minutas-api/.env (no .env.example) al arrancar."""
    env_path = Path(__file__).resolve().parent / ".env"
    if not env_path.is_file():
        return
    for raw in env_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, val = line.partition("=")
        key = key.strip()
        val = val.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = val


_cargar_env_local()


@asynccontextmanager
async def lifespan(app: FastAPI):
    repo.init_db()
    repo_import.init_import_db()
    yield


app = FastAPI(
    title="minutas-api",
    version="0.1.0",
    description="Minutas de reunión semanal — seguimiento de solicitudes a compras.",
    lifespan=lifespan,
)

_origins = os.environ.get("MINUTAS_CORS_ORIGINS", "*").strip()
_allow = ["*"] if _origins == "*" else [o.strip() for o in _origins.split(",") if o.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=_allow,
    allow_methods=["*"],
    allow_headers=["*"],
)


def _http_from_value(e: ValueError) -> HTTPException:
    return HTTPException(status.HTTP_400_BAD_REQUEST, detail=str(e))


def _http_from_lookup(e: LookupError) -> HTTPException:
    return HTTPException(status.HTTP_404_NOT_FOUND, detail=str(e))


def _email_desde_excel(sesion_id: int, asunto_custom: str | None = None):
    sesion = repo.obtener_sesion_detalle(sesion_id)
    if not sesion:
        raise LookupError("Sesión no encontrada.")
    pedidos = repo_import.listar_pedidos_agrupados(sesion_id, solo_elegibles=False)
    pedidos = [p for p in pedidos if p["seleccionada"]]
    notas = repo_import.notas_por_sesion(sesion_id)
    if pedidos:
        txt, html = construir_mail_minuta_excel(
            sesion.semana_iso,
            sesion.fecha,
            sesion.responsable,
            pedidos,
            notas,
            sesion.notas_generales,
        )
        asunto = asunto_custom or f"Minuta reunión — {sesion.semana_iso} ({sesion.fecha})"
        return asunto, txt, html
    asunto = construir_asunto(sesion, asunto_custom)
    return asunto, construir_cuerpo_texto(sesion), construir_cuerpo_html(sesion)


@app.get("/health", tags=["meta"])
def health() -> dict[str, str]:
    return {"status": "ok", "service": "minutas-api"}


@app.get("/sesiones", response_model=list[SesionResumen], tags=["sesiones"])
def listar_sesiones(limite: int = 20) -> list[SesionResumen]:
    return repo.listar_sesiones(min(limite, 100))


@app.get("/sesiones/actual", response_model=SesionDetalle | None, tags=["sesiones"])
def sesion_actual() -> SesionDetalle | None:
    sesiones = repo.listar_sesiones(50)
    for s in sesiones:
        if s.estado.value == "abierta":
            return repo.obtener_sesion_detalle(s.id)
    return None


@app.get("/sesiones/{sesion_id}", response_model=SesionDetalle, tags=["sesiones"])
def obtener_sesion(sesion_id: int) -> SesionDetalle:
    det = repo.obtener_sesion_detalle(sesion_id)
    if not det:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Sesión no encontrada.")
    return det


@app.post("/sesiones/iniciar", response_model=SesionDetalle, tags=["sesiones"])
def iniciar_sesion(body: SesionCreate) -> SesionDetalle:
    try:
        return repo.iniciar_sesion_semana(body.notas_generales, body.responsable)
    except ValueError as e:
        raise _http_from_value(e) from e


@app.patch("/sesiones/{sesion_id}", response_model=SesionResumen, tags=["sesiones"])
def patch_sesion(sesion_id: int, body: SesionUpdate) -> SesionResumen:
    try:
        return repo.actualizar_sesion(sesion_id, body.notas_generales, body.responsable)
    except LookupError as e:
        raise _http_from_lookup(e) from e
    except ValueError as e:
        raise _http_from_value(e) from e


@app.post("/sesiones/{sesion_id}/solicitudes", response_model=Solicitud, tags=["solicitudes"])
def post_solicitud(sesion_id: int, body: SolicitudCreate) -> Solicitud:
    try:
        return repo.crear_solicitud(sesion_id, body)
    except LookupError as e:
        raise _http_from_lookup(e) from e
    except ValueError as e:
        raise _http_from_value(e) from e


@app.patch("/sesiones/{sesion_id}/solicitudes/{solicitud_id}", response_model=Solicitud, tags=["solicitudes"])
def patch_solicitud(sesion_id: int, solicitud_id: int, body: SolicitudUpdate) -> Solicitud:
    try:
        return repo.actualizar_solicitud(sesion_id, solicitud_id, body)
    except LookupError as e:
        raise _http_from_lookup(e) from e
    except ValueError as e:
        raise _http_from_value(e) from e


@app.post(
    "/sesiones/{sesion_id}/solicitudes/{solicitud_id}/entregas",
    response_model=Solicitud,
    tags=["solicitudes"],
)
def post_entrega(sesion_id: int, solicitud_id: int, body: EntregaParcialCreate) -> Solicitud:
    try:
        return repo.agregar_entrega(sesion_id, solicitud_id, body)
    except LookupError as e:
        raise _http_from_lookup(e) from e
    except ValueError as e:
        raise _http_from_value(e) from e


@app.post(
    "/sesiones/{sesion_id}/solicitudes/{solicitud_id}/actualizaciones",
    response_model=Actualizacion,
    tags=["solicitudes"],
)
def post_actualizacion(
    sesion_id: int, solicitud_id: int, body: ActualizacionCreate
) -> Actualizacion:
    try:
        return repo.agregar_actualizacion_solicitud(sesion_id, solicitud_id, body)
    except LookupError as e:
        raise _http_from_lookup(e) from e


@app.post("/sesiones/{sesion_id}/temas", response_model=Tema, tags=["temas"])
def post_tema(sesion_id: int, body: TemaCreate) -> Tema:
    try:
        return repo.crear_tema(sesion_id, body)
    except LookupError as e:
        raise _http_from_lookup(e) from e
    except ValueError as e:
        raise _http_from_value(e) from e


@app.patch("/temas/{tema_id}", response_model=Tema, tags=["temas"])
def patch_tema(tema_id: int, body: TemaUpdate) -> Tema:
    try:
        return repo.actualizar_tema(tema_id, body)
    except LookupError as e:
        raise _http_from_lookup(e) from e


@app.post("/sesiones/{sesion_id}/abandonar", response_model=SesionResumen, tags=["sesiones"])
def abandonar_sesion(sesion_id: int) -> SesionResumen:
    try:
        return repo.cerrar_sesion_sin_enviar(sesion_id)
    except LookupError as e:
        raise _http_from_lookup(e) from e
    except ValueError as e:
        raise _http_from_value(e) from e


@app.post("/sesiones/{sesion_id}/importar-excel", tags=["import"])
async def importar_excel(sesion_id: int, archivo: UploadFile = File(...)):
    if not archivo.filename or not archivo.filename.lower().endswith((".xlsx", ".xlsm", ".xls")):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Formato no soportado. Use .xlsx")
    try:
        raw = await archivo.read()
        # parsear_excel es CPU/IO pesado: fuera del event loop para no colgar /health ni la carga inicial
        hoja, filas = await asyncio.to_thread(parsear_excel, raw)
        meta = await asyncio.to_thread(
            repo_import.guardar_importacion,
            sesion_id,
            archivo.filename,
            hoja,
            filas,
            False,
            raw,
        )
        return meta
    except LookupError as e:
        raise _http_from_lookup(e) from e
    except ValueError as e:
        raise _http_from_value(e) from e
    except Exception as e:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail=f"Error leyendo Excel: {e}") from e


@app.get("/sesiones/{sesion_id}/pedidos", tags=["import"])
def listar_pedidos(sesion_id: int, solo_elegibles: bool = True):
    return repo_import.listar_pedidos_agrupados(sesion_id, solo_elegibles=solo_elegibles)


@app.get("/sesiones/{sesion_id}/exportar-excel", tags=["export"])
def exportar_excel_endpoint(
    sesion_id: int,
    refs: list[str] | None = Query(None),
    cols: list[str] | None = Query(None),
):
    sesion = repo.obtener_sesion_detalle(sesion_id)
    if not sesion:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Sesión no encontrada.")
    if sesion.estado.value == "enviada":
        raise HTTPException(status.HTTP_403_FORBIDDEN, detail="No se puede exportar una sesión enviada.")

    source_bytes = repo_import.obtener_source_bytes(sesion_id)
    filas = repo_import.listar_filas(sesion_id, order_by_fila_excel=True)
    if refs:
        filas = [f for f in filas if f["ref_pedido"] in refs]
    selected_refs = [f["ref_pedido"] for f in filas] if filas else None

    # Si no hay source_bytes (sesión vieja), generar desde filas DB
    if source_bytes is None and filas:
        output = exportar_excel(None, selected_refs=selected_refs, cols=cols, filas_db=filas)
    elif source_bytes is None:
        raise HTTPException(
            status.HTTP_404_NOT_FOUND, detail="No se encontró archivo de importación para esta sesión."
        )
    else:
        # Si tenemos filas filtradas, pasar filas_db para evitar re-parsear todo el archivo
        if filas is not None:
            output = exportar_excel(source_bytes, selected_refs=selected_refs, cols=cols, filas_db=filas)
        else:
            output = exportar_excel(source_bytes, selected_refs=selected_refs, cols=cols)

    filename = f"Minuta_{sesion.semana_iso}_{sesion.fecha}.xlsx"
    # Sanitize filename for header
    safe = "".join(c if c.isalnum() or c in "-_." else "_" for c in filename)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{safe}"'},
    )


@app.post("/sesiones/{sesion_id}/importar-novedades", tags=["import"])
async def importar_novedades(sesion_id: int, archivo: UploadFile = File(...)):
    sesion = repo.obtener_sesion_detalle(sesion_id)
    if not sesion:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Sesión no encontrada.")
    if sesion.estado.value == "enviada":
        raise HTTPException(
            status.HTTP_403_FORBIDDEN, detail="No se puede importar novedades en una sesión enviada."
        )
    if not archivo.filename or not archivo.filename.lower().endswith((".xlsx", ".xlsm", ".xls")):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Formato no soportado. Use .xlsx")

    raw = await archivo.read()

    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        if "Novedades" not in wb.sheetnames:
            wb.close()
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST, detail="No se encontró la hoja 'Novedades' en el archivo."
            )

        ws = wb["Novedades"]
        # Leer encabezados para mapear columnas por nombre (soporta export secuencial o con gaps)
        header_row = None
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            header_row = row
            break
        if header_row is None:
            wb.close()
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Hoja 'Novedades' vacía.")

        # Normalizar headers a índice
        header_map: dict[str, int] = {}
        for idx, h in enumerate(header_row):
            if h:
                header_map[str(h).strip().lower()] = idx

        # Índices por nombre, con fallback a posiciones secuenciales del export
        idx_novedades = header_map.get("novedades", header_map.get("novedad", header_map.get("actualización reunión", 22)))
        idx_num_odoo = header_map.get("num_odoo", header_map.get("num odoo", 5))
        idx_num_solicitud = header_map.get("num_solicitud", header_map.get("num solicitud", 7))
        # fila_excel oculto está en AA/header "fila_excel" (col 27 -> idx 26)
        idx_fila_excel = header_map.get("fila_excel", header_map.get("fila", 26))

        filas: list[FilaPedido] = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            def get_cell(idx: int | None) -> str:
                if idx is None or idx >= len(row):
                    return ""
                v = row[idx]
                return str(v).strip() if v is not None else ""

            num_odoo = get_cell(idx_num_odoo if isinstance(idx_num_odoo, int) else None)
            num_solicitud = get_cell(idx_num_solicitud if isinstance(idx_num_solicitud, int) else None)
            novedades = get_cell(idx_novedades if isinstance(idx_novedades, int) else None)
            raw_fila_excel_val = row[idx_fila_excel] if isinstance(idx_fila_excel, int) and idx_fila_excel < len(row) else None
            try:
                fila_excel = int(raw_fila_excel_val) if raw_fila_excel_val is not None and str(raw_fila_excel_val).strip() != "" else 0
            except (ValueError, TypeError):
                fila_excel = 0

            ref_pedido = _ref_pedido(num_odoo, num_solicitud)
            if not ref_pedido and not novedades.strip():
                continue

            filas.append(
                FilaPedido(
                    fila_excel=fila_excel,
                    cant_articulos_pedido="",
                    solicitante="",
                    tipo_solicitud="",
                    maquina_linea="",
                    fecha_solicitud="",
                    ref_pedido=ref_pedido,
                    num_odoo=num_odoo,
                    num_solicitud=num_solicitud,
                    almacenista="",
                    codigo="",
                    descripcion="",
                    cantidad="",
                    unidad="",
                    precio="",
                    total="",
                    moneda="",
                    proveedor="",
                    oc_rq="",
                    fecha_oc="",
                    comprador="",
                    fecha_envio_compras="",
                    estado_item="",
                    estado_solicitud="",
                    cumplida=False,
                    elegible_reunion=True,
                    novedades=novedades,
                )
            )

        wb.close()

        if not filas:
            return {
                "procesadas": 0,
                "omitidas_duplicadas": 0,
                "pendientes_consulta": 0,
                "no_reconocidas": 0,
                "total_importadas": 0,
            }

        meta = await asyncio.to_thread(
            repo_import.guardar_importacion,
            sesion_id,
            archivo.filename,
            "Novedades",
            filas,
            True,
            raw,
        )
        return meta
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail=f"Error leyendo Novedades: {e}") from e


@app.put("/sesiones/{sesion_id}/seleccion-reunion", tags=["import"])
def seleccion_reunion(sesion_id: int, body: dict):
    refs = body.get("refs") or []
    if not isinstance(refs, list):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="refs debe ser una lista")
    n = repo_import.fijar_seleccion_reunion(sesion_id, [str(r) for r in refs])
    return {"seleccionadas": n, "refs": refs}


@app.patch("/sesiones/{sesion_id}/seleccion-pedidos", tags=["import"])
def seleccion_pedidos(sesion_id: int, body: dict):
    refs = body.get("refs") or []
    seleccionada = bool(body.get("seleccionada", True))
    if not isinstance(refs, list):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="refs debe ser una lista")
    n = repo_import.actualizar_seleccion(sesion_id, [str(r) for r in refs], seleccionada)
    return {"actualizadas": n}


@app.post("/sesiones/{sesion_id}/filas/{fila_id}/notas", tags=["import"])
def nota_fila(sesion_id: int, fila_id: int, body: ActualizacionCreate):
    try:
        return repo_import.agregar_nota_fila(sesion_id, fila_id, body.texto, body.autor)
    except LookupError as e:
        raise _http_from_lookup(e) from e


@app.get("/sesiones/{sesion_id}/consultas-pendientes", tags=["import"])
def listar_consultas(sesion_id: int):
    sesion = repo.obtener_sesion_detalle(sesion_id)
    if not sesion:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Sesión no encontrada.")
    return repo_import.listar_consultas_pendientes(sesion_id)


@app.get("/sesiones/{sesion_id}/preview-email", response_model=PreviewEmailResponse, tags=["email"])
def preview_email(sesion_id: int, asunto: str | None = None) -> PreviewEmailResponse:
    try:
        a, txt, html = _email_desde_excel(sesion_id, asunto)
    except LookupError as e:
        raise _http_from_lookup(e) from e
    return PreviewEmailResponse(asunto=a, cuerpo_texto=txt, cuerpo_html=html)


@app.post("/sesiones/{sesion_id}/enviar", response_model=EnviarMinutaResponse, tags=["email"])
def enviar_minuta_endpoint(sesion_id: int, body: EnviarMinutaRequest) -> EnviarMinutaResponse:
    try:
        destinatarios = destinatarios_default_o([str(e) for e in body.destinatarios])
        asunto, texto, html = _email_desde_excel(sesion_id, body.asunto)
        enviar_minuta(destinatarios, asunto, texto, html)
        repo.marcar_sesion_enviada(sesion_id)
    except LookupError as e:
        raise _http_from_lookup(e) from e
    except EmailNoConfigurado as e:
        raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, detail=str(e)) from e
    except ValueError as e:
        raise _http_from_value(e) from e
    except smtplib.SMTPException as e:
        raise HTTPException(
            status.HTTP_502_BAD_GATEWAY,
            detail=f"Error al enviar correo: {e}",
        ) from e

    return EnviarMinutaResponse(
        ok=True,
        mensaje="Minuta enviada correctamente.",
        sesion_id=sesion_id,
        destinatarios=destinatarios,
    )
